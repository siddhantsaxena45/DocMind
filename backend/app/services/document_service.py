import io
import re
import base64
import zipfile
import xml.etree.ElementTree as ET
from typing import List, Tuple, Dict, Any, Optional
from pypdf import PdfReader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.messages import HumanMessage
from langchain_google_genai import ChatGoogleGenerativeAI
from app.core.config import settings, api_key_rotator


URL_REGEX = re.compile(
    r'(?:https?://[^\s<>"\'\(\)]+|www\.[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}[^\s<>"\'\(\)]*|[a-zA-Z0-9\-]+(?:\.[a-zA-Z0-9\-]+)*\.(?:ai|com|org|io|dev|app|net|co|in|edu|gov|me|xyz)(?:/[^\s<>"\'\(\)]*)?|(?:github\.com|linkedin\.com/in|leetcode\.com|kaggle\.com|huggingface\.co)/[a-zA-Z0-9_.\-]+(?:/[a-zA-Z0-9_.\-]+)*)',
    re.IGNORECASE
)

MARKDOWN_LINK_REGEX = re.compile(
    r'\[([^\]]+)\]\((https?://[^\s)]+)\)'
)


def clean_url(url: Optional[str]) -> Optional[str]:
    if not url:
        return None
    u = url.strip().rstrip(".,;!?:)'\"]}>")
    if not u:
        return None
    if u.startswith("www."):
        u = "https://" + u
    elif re.search(r'^[a-zA-Z0-9\-]+(?:\.[a-zA-Z0-9\-]+)*\.(?:ai|com|org|io|dev|app|net|co|in|edu|gov|me|xyz)(?:/.*)?$', u, re.IGNORECASE):
        u = "https://" + u
    elif any(u.lower().startswith(prefix) for prefix in ["github.com", "linkedin.com", "leetcode.com", "kaggle.com", "huggingface.co"]):
        u = "https://" + u
    if not (u.startswith("http://") or u.startswith("https://") or u.startswith("mailto:")):
        return None
    return u


def extract_urls_from_text(text: str) -> List[str]:
    """
    Extracts Markdown links, HTTP/HTTPS URLs, www links, and platform handles (GitHub, LinkedIn, LeetCode)
    from raw or formatted text.
    """
    found: List[str] = []
    if not text:
        return found

    for match in MARKDOWN_LINK_REGEX.finditer(text):
        target = clean_url(match.group(2))
        if target and target not in found:
            found.append(target)

    for match in URL_REGEX.finditer(text):
        target = clean_url(match.group(0))
        if target and target not in found:
            found.append(target)

    return found


class DocumentParseResult:
    def __init__(self, full_text: str, pages_data: List[Dict[str, Any]], links: List[str]):
        self.full_text = full_text
        self.pages_data = pages_data  # list of {"page": int, "text": str}
        self.links = links


class DocumentService:
    def __init__(self):
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=800,
            chunk_overlap=120,
            separators=["\n\n", "\n", ". ", " ", ""]
        )

    def extract_from_pdf(self, content: bytes) -> DocumentParseResult:
        pdf_reader = PdfReader(io.BytesIO(content))
        pages_data = []
        links: List[str] = []
        full_text_parts = []

        for idx, page in enumerate(pdf_reader.pages, start=1):
            page_text = page.extract_text() or ""
            page_links: List[str] = []

            # 1. Extract PDF Annotations (/Link -> /A -> /URI)
            annots = getattr(page, "annotations", None)
            if annots is None and "/Annots" in page:
                annots_obj = page["/Annots"]
                annots = annots_obj.get_object() if hasattr(annots_obj, "get_object") else annots_obj

            if annots:
                for annot in annots:
                    obj = annot.get_object() if hasattr(annot, "get_object") else annot
                    if isinstance(obj, dict) and obj.get("/Subtype") == "/Link":
                        action = obj.get("/A")
                        if action:
                            action = action.get_object() if hasattr(action, "get_object") else action
                        if isinstance(action, dict) and "/URI" in action:
                            uri = action["/URI"]
                            uri = uri.get_object() if hasattr(uri, "get_object") else uri
                            if isinstance(uri, bytes):
                                try:
                                    uri = uri.decode("utf-8", errors="ignore")
                                except Exception:
                                    uri = None
                            c_uri = clean_url(uri)
                            if c_uri and c_uri not in page_links:
                                page_links.append(c_uri)

            # 2. Extract plaintext URLs from page text
            for u in extract_urls_from_text(page_text):
                if u not in page_links:
                    page_links.append(u)

            clean_text = page_text.strip()
            # 3. Inline page hyperlinks directly into the page text so RAG chunking preserves them
            if page_links:
                clean_text += f"\n\n[Page {idx} Embedded Links:\n" + "\n".join(f"- {l}" for l in page_links) + "]"
                for l in page_links:
                    if l not in links:
                        links.append(l)

            if clean_text:
                pages_data.append({"page": idx, "text": clean_text})
                full_text_parts.append(f"--- Page {idx} ---\n{clean_text}")

        full_text = "\n\n".join(full_text_parts)
        if links:
            full_text += "\n\n[DOCUMENT VERIFIED HYPERLINKS]\n" + "\n".join(links)
        return DocumentParseResult(full_text=full_text, pages_data=pages_data, links=links)

    def extract_from_docx(self, content: bytes) -> DocumentParseResult:
        links: List[str] = []
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))

            # Build map of relationship ID -> target URL
            rel_hyperlinks: Dict[str, str] = {}
            for rel_id, rel in doc.part.rels.items():
                target = getattr(rel, "target_ref", "")
                if target and (getattr(rel, "is_external", False) or "hyperlink" in getattr(rel, "reltype", "").lower() or str(target).startswith("http")):
                    c_target = clean_url(str(target))
                    if c_target:
                        rel_hyperlinks[rel_id] = c_target

            paragraphs: List[str] = []
            for p in doc.paragraphs:
                p_parts: List[str] = []
                for child in p._element:
                    tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
                    if tag == "hyperlink":
                        r_id = child.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id") or child.attrib.get("r:id")
                        target_url = rel_hyperlinks.get(r_id) if r_id else None
                        h_text = "".join(node.text for node in child.iter() if node.tag.endswith("t") and node.text).strip()
                        if target_url and target_url not in links:
                            links.append(target_url)
                        if target_url and h_text:
                            p_parts.append(f"{h_text} ({target_url})")
                        elif h_text:
                            p_parts.append(h_text)
                        elif target_url:
                            p_parts.append(target_url)
                    elif tag == "r":
                        r_text = "".join(node.text for node in child.iter() if node.tag.endswith("t") and node.text)
                        if r_text:
                            p_parts.append(r_text)

                para_str = "".join(p_parts).strip() if p_parts else p.text.strip()
                if para_str:
                    paragraphs.append(para_str)

            # Extract table cells
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        paragraphs.append(" | ".join(row_cells))

            text = "\n\n".join(paragraphs)
        except Exception:
            # Fallback to direct XML parse
            text, xml_links = self._extract_docx_xml(content)
            for xl in xml_links:
                if xl not in links:
                    links.append(xl)

        # Extract plaintext URLs from docx text
        for u in extract_urls_from_text(text):
            if u not in links:
                links.append(u)

        if links:
            text += "\n\n[Document Embedded Links:\n" + "\n".join(f"- {l}" for l in links) + "]"

        pages_data = [{"page": 1, "text": text}]
        return DocumentParseResult(full_text=text, pages_data=pages_data, links=links)

    def _extract_docx_xml(self, content: bytes) -> Tuple[str, List[str]]:
        xml_links: List[str] = []
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                # 1. Parse relationships for hyperlinks
                rel_map: Dict[str, str] = {}
                if "word/_rels/document.xml.rels" in z.namelist():
                    rels_data = z.read("word/_rels/document.xml.rels")
                    rels_tree = ET.fromstring(rels_data)
                    for rel in rels_tree:
                        t = rel.attrib.get("Target", "")
                        r_id = rel.attrib.get("Id", "")
                        c_t = clean_url(t)
                        if c_t and r_id:
                            rel_map[r_id] = c_t
                            if c_t not in xml_links:
                                xml_links.append(c_t)

                # 2. Parse document XML
                xml_content = z.read("word/document.xml")
                tree = ET.fromstring(xml_content)
                namespaces = {
                    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
                }
                texts = [node.text for node in tree.findall(".//w:t", namespaces) if node.text]
                return " ".join(texts), xml_links
        except Exception as e:
            return f"Error extracting DOCX text: {e}", xml_links

    def extract_from_pptx(self, content: bytes) -> DocumentParseResult:
        pages_data = []
        full_text_parts = []
        links: List[str] = []

        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                slide_files = sorted([f for f in z.namelist() if f.startswith("ppt/slides/slide") and f.endswith(".xml")])
                for idx, slide_path in enumerate(slide_files, start=1):
                    xml_data = z.read(slide_path)
                    tree = ET.fromstring(xml_data)
                    namespaces = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                    slide_texts = [node.text for node in tree.findall(".//a:t", namespaces) if node.text]
                    slide_content = " ".join(slide_texts).strip()

                    slide_links: List[str] = []
                    # Check slide relationship file for embedded hyperlinks
                    rel_path = slide_path.replace("ppt/slides/slide", "ppt/slides/_rels/slide") + ".rels"
                    if rel_path in z.namelist():
                        try:
                            rels_xml = z.read(rel_path)
                            tree_rels = ET.fromstring(rels_xml)
                            for rel in tree_rels:
                                target = rel.attrib.get("Target", "")
                                rel_type = rel.attrib.get("Type", "")
                                target_mode = rel.attrib.get("TargetMode", "")
                                if target and (target_mode == "External" or "hyperlink" in rel_type.lower() or target.startswith("http")):
                                    c_target = clean_url(target)
                                    if c_target and c_target not in slide_links:
                                        slide_links.append(c_target)
                        except Exception:
                            pass

                    # Plaintext URLs in slide content
                    for u in extract_urls_from_text(slide_content):
                        if u not in slide_links:
                            slide_links.append(u)

                    if slide_links:
                        slide_content += f"\n\n[Slide {idx} Links:\n" + "\n".join(f"- {l}" for l in slide_links) + "]"
                        for l in slide_links:
                            if l not in links:
                                links.append(l)

                    if slide_content:
                        pages_data.append({"page": idx, "text": slide_content})
                        full_text_parts.append(f"--- Slide {idx} ---\n{slide_content}")
        except Exception as e:
            full_text_parts.append(f"Error parsing PPTX: {e}")

        full_text = "\n\n".join(full_text_parts)
        if links:
            full_text += "\n\n[DOCUMENT VERIFIED HYPERLINKS]\n" + "\n".join(links)
        if not pages_data:
            pages_data = [{"page": 1, "text": full_text or "No text found in presentation."}]
        return DocumentParseResult(full_text=full_text, pages_data=pages_data, links=links)

    def extract_from_xlsx(self, content: bytes) -> DocumentParseResult:
        pages_data = []
        full_text_parts = []
        links: List[str] = []

        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            for idx, sheet_name in enumerate(wb.sheetnames, start=1):
                sheet = wb[sheet_name]
                rows_text = []
                sheet_links: List[str] = []

                for row in sheet.iter_rows():
                    row_vals = []
                    for cell in row:
                        val = str(cell.value).strip() if cell.value is not None else ""
                        target = None
                        if hasattr(cell, "hyperlink") and cell.hyperlink:
                            target = getattr(cell.hyperlink, "target", None)
                        if target:
                            c_target = clean_url(str(target))
                            if c_target:
                                val = f"{val} ({c_target})" if val else c_target
                                if c_target not in sheet_links:
                                    sheet_links.append(c_target)
                        if val:
                            row_vals.append(val)
                    if row_vals:
                        rows_text.append(" | ".join(row_vals))

                sheet_text = f"Sheet: {sheet_name}\n" + "\n".join(rows_text)

                for u in extract_urls_from_text(sheet_text):
                    if u not in sheet_links:
                        sheet_links.append(u)

                if sheet_links:
                    sheet_text += f"\n\n[Sheet {sheet_name} Links:\n" + "\n".join(f"- {l}" for l in sheet_links) + "]"
                    for l in sheet_links:
                        if l not in links:
                            links.append(l)

                pages_data.append({"page": idx, "text": sheet_text})
                full_text_parts.append(f"--- Sheet {idx} ({sheet_name}) ---\n{sheet_text}")
        except Exception as e:
            full_text_parts.append(f"Error parsing Excel spreadsheet: {e}")

        full_text = "\n\n".join(full_text_parts)
        if links:
            full_text += "\n\n[DOCUMENT VERIFIED HYPERLINKS]\n" + "\n".join(links)
        if not pages_data:
            pages_data = [{"page": 1, "text": full_text or "No table data found."}]
        return DocumentParseResult(full_text=full_text, pages_data=pages_data, links=links)

    def extract_from_image(self, content: bytes, filename: str) -> DocumentParseResult:
        keys = api_key_rotator.get_rotated_google_keys()
        if not keys:
            return DocumentParseResult(
                full_text="[Image text extraction unavailable: No Google API key configured]",
                pages_data=[{"page": 1, "text": "[Image: " + filename + "]"}],
                links=[]
            )

        mime_type = "image/png" if filename.lower().endswith(".png") else "image/jpeg"
        b64_image = base64.b64encode(content).decode("utf-8")

        prompt = (
            "You are an expert OCR and document analysis system. "
            "Transcribe all text, numbers, formulas, tables, labels, and diagrams from this document/image. "
            "Output clear, structured markdown. Do not include introductory conversational filler."
        )

        ocr_text = ""
        for key in keys:
            try:
                llm = ChatGoogleGenerativeAI(
                    model=settings.GEMINI_MODEL,
                    google_api_key=key,
                    temperature=0.0
                )
                message = HumanMessage(
                    content=[
                        {"type": "text", "text": prompt},
                        {
                            "type": "image_url",
                            "image_url": {"url": f"data:{mime_type};base64,{b64_image}"}
                        }
                    ]
                )
                res = llm.invoke([message])
                ocr_text = res.content if isinstance(res.content, str) else str(res.content)
                if ocr_text.strip():
                    break
            except Exception as e:
                print(f"Image OCR attempt with key failed: {e}")
                continue

        if not ocr_text.strip():
            ocr_text = f"[Image: {filename} - OCR could not recognize text]"

        links = extract_urls_from_text(ocr_text)
        if links:
            ocr_text += "\n\n[Image Extracted Links:\n" + "\n".join(f"- {l}" for l in links) + "]"

        pages_data = [{"page": 1, "text": ocr_text}]
        return DocumentParseResult(full_text=ocr_text, pages_data=pages_data, links=links)

    def extract_from_text(self, content: bytes) -> DocumentParseResult:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception:
                text = content.decode("utf-8", errors="ignore")

        links = extract_urls_from_text(text)
        if links:
            text += "\n\n[Document Embedded Links:\n" + "\n".join(f"- {l}" for l in links) + "]"

        pages_data = [{"page": 1, "text": text}]
        return DocumentParseResult(full_text=text, pages_data=pages_data, links=links)

    def parse_document(self, content: bytes, filename: str) -> DocumentParseResult:
        ext = filename.lower().split(".")[-1] if "." in filename else ""
        if ext == "pdf":
            return self.extract_from_pdf(content)
        elif ext in ["docx", "doc"]:
            return self.extract_from_docx(content)
        elif ext in ["pptx", "ppt"]:
            return self.extract_from_pptx(content)
        elif ext in ["xlsx", "xls", "csv"]:
            if ext == "csv":
                return self.extract_from_text(content)
            return self.extract_from_xlsx(content)
        elif ext in ["png", "jpg", "jpeg", "webp", "bmp"]:
            return self.extract_from_image(content, filename)
        else:
            # Default to text / markdown / code
            return self.extract_from_text(content)

    def chunk_document(
        self,
        parse_result: DocumentParseResult,
        user_id: str,
        filename: str,
        document_id: str
    ) -> List[Document]:
        all_chunks: List[Document] = []
        chunk_idx = 0

        # Chunk per page to retain accurate page citations
        for page_info in parse_result.pages_data:
            page_num = page_info.get("page", 1)
            page_text = page_info.get("text", "").strip()
            if not page_text:
                continue

            page_chunks = self.splitter.split_text(page_text)
            for chunk_content in page_chunks:
                chunk_idx += 1
                metadata = {
                    "user_id": user_id,
                    "document_id": document_id,
                    "filename": filename,
                    "page": page_num,
                    "chunk_id": f"{document_id}_{chunk_idx}"
                }
                all_chunks.append(Document(page_content=chunk_content, metadata=metadata))

        # Fallback if pages were somehow empty but full_text exists
        if not all_chunks and parse_result.full_text.strip():
            raw_chunks = self.splitter.split_text(parse_result.full_text)
            for c in raw_chunks:
                chunk_idx += 1
                all_chunks.append(
                    Document(
                        page_content=c,
                        metadata={
                            "user_id": user_id,
                            "document_id": document_id,
                            "filename": filename,
                            "page": 1,
                            "chunk_id": f"{document_id}_{chunk_idx}"
                        }
                    )
                )

        return all_chunks

document_service = DocumentService()
